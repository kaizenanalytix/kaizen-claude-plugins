#!/usr/bin/env python3
"""
build_scorecard.py — Kaizen Project Governance Scorecard generator.

Two subcommands, mirroring the split used across the PDP skills (script = deterministic
work, model = judgment):

  scan   Walk the project's PDP phase folders and emit a JSON pre-fill: which PDP-checklist
         items have a file on disk (with last-reviewed date + path), which are missing (gaps),
         recent files as milestone candidates, and — if a prior scorecard deck exists — the
         previous month's status notes for carry-forward. The MODEL reads this, asks the EL
         about gaps and the subjective status blocks, and writes the data JSON.

  build  Populate the branded Project Governance Scorecard slide from a data JSON and save it
         to "7. Project Governance". If a scorecard deck already exists there, a NEW dated
         slide is appended (cloned from the bundled blank template slide); otherwise a new deck
         is created (title slide + this month's slide). Never overwrites an existing deck's
         earlier slides (G1-friendly: history accretes).

The status RAG light is rendered as a fresh oval placed at the computed centre of each status
row (the template's own ovals are inconsistently positioned, so we drop clean ones instead).

Data JSON schema (see build):
{
  "project_id": "TMNA FY2027 Trade Vault",
  "date": "Jul 2026",
  "product_name": "Trade Vault",              # optional banner text; defaults to project_id
  "status": {
    "Scope":       {"rag": "green", "notes": "..."},
    "Timeline":    {"rag": "amber", "notes": "..."},
    "Budget":      {"rag": "red",   "notes": "..."},
    "Client Sat.": {"rag": "green", "notes": "..."},
    "Kaizen Team": {"rag": "green", "notes": "..."},
    "Pursuits":    {"rag": "green", "notes": "..."}
  },
  "milestones": [ {"milestone": "...", "date": "1/26", "notes": "..."}, ... ],
  "pdp_checklist": [ {"item": "RAID Log", "last_reviewed": "Jul 2026", "notes": "..."}, ... ]
}
"""

import argparse, copy, datetime as dt, json, os, re, sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- constants ----
RAG = {
    "green": RGBColor(0x92, 0xD0, 0x50),
    "amber": RGBColor(0xFF, 0xFF, 0x00),
    "red":   RGBColor(0xFF, 0x00, 0x00),
    "grey":  RGBColor(0xBF, 0xBF, 0xBF),
}
RAG_ALIASES = {
    "g": "green", "green": "green", "on track": "green", "ontrack": "green", "ok": "green",
    "a": "amber", "amber": "amber", "yellow": "amber", "at risk": "amber", "atrisk": "amber", "watch": "amber",
    "r": "red", "red": "red", "off track": "red", "offtrack": "red", "blocked": "red",
    "": "grey", "none": "grey", "grey": "grey", "gray": "grey", "n/a": "grey", "na": "grey",
}

STATUS_ROWS = ["Scope", "Timeline", "Budget", "Client Sat.", "Kaizen Team", "Pursuits"]

# PDP-checklist item -> filename signals (case-insensitive substrings; any group hit = match).
# This is the explicit alias bridge so detection does not depend on fuzzy name overlap.
CHECKLIST_ALIASES = {
    "RAID Log":                 ["raid"],
    "Business Continuity Plan": ["business continuity", "bcp", "continuity plan"],
    "End to End Diagram":       ["e2e", "end to end", "end-to-end", "solution overview", "playbook", "process flow"],
    "System Arch Diagram":      ["system arch", "architecture", "solution architecture", "technical design", "tech design", "tdd"],
    "Value Tracker":            ["value tracker", "benefit tracker", "benefits tracker", "value track", "p3", "economics", "roi"],
    "Case Study":               ["case study", "casestudy"],
    # extra items that appear in some historical scorecards
    "Technical Solution Architecture": ["architecture", "solution architecture", "technical design", "tdd", "tech design"],
    "JIRA Board":               ["jira"],
}

PHASE_FOLDERS = [
    "0. Sales Alignment", "1. Sales Handoff & Transition", "2. Plan", "3. Analyze",
    "4. Design", "5. Develop", "6. Deploy", "7. Project Governance", "8. Quality",
]
GOV_FOLDER = "7. Project Governance"
SKIP_DIRS = {".git", "__pycache__", ".ipynb_checkpoints", ".DS_Store"}
TRACKER_HINTS = ("report card", "phase checklist", "placement & naming", "audit trail",
                 "governance scorecard")  # never treat trackers/this deck as deliverables

TEMPLATE = Path(__file__).resolve().parent.parent / "template" / "Project Governance Scorecard - TEMPLATE.pptx"
BLANK_SLIDE_INDEX = 3   # the empty scorecard slide in the bundled template


# ============================================================ helpers: scan ====
def _iter_files(root: Path):
    for phase in PHASE_FOLDERS:
        folder = root / phase
        if not folder.is_dir():
            continue
        for dp, dns, fns in os.walk(folder):
            dns[:] = [d for d in dns if d not in SKIP_DIRS]
            for fn in fns:
                if fn.startswith("~$") or fn.startswith("."):
                    continue
                low = fn.lower()
                if any(h in low for h in TRACKER_HINTS):
                    continue
                p = Path(dp) / fn
                yield phase, p, low


def _fmt_month(ts: float) -> str:
    return dt.datetime.fromtimestamp(ts).strftime("%b %Y")


def scan(root: Path, project_id: str, output_deck: Path):
    files = list(_iter_files(root))

    # --- checklist detection via alias substrings; newest matching file wins ---
    checklist = []
    for item, needles in CHECKLIST_ALIASES.items():
        hits = []
        for phase, p, low in files:
            if any(n in low for n in needles):
                hits.append(p)
        if hits:
            newest = max(hits, key=lambda x: x.stat().st_mtime)
            checklist.append({
                "item": item, "found": True,
                "last_reviewed": _fmt_month(newest.stat().st_mtime),
                "file": newest.name,
                "path": str(newest.relative_to(root)).replace("\\", "/"),
                "candidates": [str(h.relative_to(root)).replace("\\", "/") for h in hits],
            })
        else:
            checklist.append({"item": item, "found": False, "last_reviewed": "", "file": "", "path": "", "candidates": []})

    gaps = [c["item"] for c in checklist if not c["found"]]

    # --- milestone candidates: files touched in the last 45 days, newest first ---
    cutoff = dt.datetime.now().timestamp() - 45 * 86400
    recent = sorted(
        [(p, p.stat().st_mtime, phase) for phase, p, _ in files if p.stat().st_mtime >= cutoff],
        key=lambda x: x[1], reverse=True)[:12]
    milestone_candidates = [{
        "name": p.stem, "date": dt.datetime.fromtimestamp(m).strftime("%-m/%-d") if hasattr(dt.datetime, 'now') else dt.datetime.fromtimestamp(m).strftime("%m/%d"),
        "phase": phase, "path": str(p.relative_to(root)).replace("\\", "/"),
    } for p, m, phase in recent]

    # --- carry-forward: previous month's status notes, if a deck already exists ---
    status_prefill, prior_month = {}, ""
    if output_deck.exists():
        try:
            status_prefill, prior_month = _read_last_status(output_deck)
        except Exception as e:
            status_prefill = {"_error": f"could not read prior deck: {e}"}

    return {
        "project_id": project_id,
        "root": str(root),
        "output_deck": str(output_deck),
        "deck_exists": output_deck.exists(),
        "prior_month": prior_month,
        "checklist": checklist,
        "gaps": gaps,
        "status_rows_needed": STATUS_ROWS,
        "status_prefill": status_prefill,
        "milestone_candidates": milestone_candidates,
    }


# ================================================ helpers: pptx table access ====
def _tables(slide):
    return [sh for sh in slide.shapes if sh.has_table]


def _find_table(slide, header0):
    for sh in _tables(slide):
        if sh.table.cell(0, 0).text.strip().lower() == header0.lower():
            return sh
    return None


def _set_cell(cell, text, size=9):
    cell.text = "" if text is None else str(text)
    for para in cell.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size)


def _read_last_status(deck_path: Path):
    """Best-effort read of the most recent slide's status notes for carry-forward."""
    prs = Presentation(str(deck_path))
    for slide in reversed(prs.slides):
        st = _find_table(slide, "Status")
        if st is None:
            continue
        tbl = st.table
        out = {}
        for r in range(1, len(tbl.rows)):
            label = tbl.cell(r, 0).text.strip()
            notes = tbl.cell(r, 2).text.strip()
            if label:
                out[label] = {"notes": notes}
        # month from the date placeholder if present
        month = ""
        for sh in slide.shapes:
            if sh.has_text_frame and re.search(r"\b(20\d\d)\b", sh.text_frame.text or ""):
                month = sh.text_frame.text.strip()
                break
        return out, month
    return {}, ""


# ==================================================== helpers: slide cloning ====
def _clone_blank_slide(out_prs):
    """Append a copy of the bundled blank scorecard slide into out_prs and return it."""
    tmpl = Presentation(str(TEMPLATE))
    blank = tmpl.slides[BLANK_SLIDE_INDEX]
    layout_name = blank.slide_layout.name

    # find a matching layout in the output deck (it was built from the same template)
    layout = None
    for master in out_prs.slide_masters:
        for lo in master.slide_layouts:
            if lo.name == layout_name:
                layout = lo
                break
        if layout:
            break
    if layout is None:
        layout = out_prs.slide_layouts[0]

    new_slide = out_prs.slides.add_slide(layout)
    # strip auto-added placeholders from the layout
    for shp in list(new_slide.shapes):
        shp._element.getparent().remove(shp._element)
    # deep-copy every shape element from the blank template slide (no media rels in this template)
    for shp in blank.shapes:
        new_slide.shapes._spTree.append(copy.deepcopy(shp._element))
    return new_slide


def _delete_slides(prs, indexes):
    """Fully remove slides by index: drop the presentation->slide relationship (so the
    part is no longer serialized) and remove the sldId entry. Prevents duplicate-partname
    collisions when new slides are later added."""
    sldIdLst = prs.slides._sldIdLst
    sldIds = list(sldIdLst)
    for i in sorted(indexes, reverse=True):
        if i >= len(sldIds):
            continue
        sldId = sldIds[i]
        rId = sldId.get(qn('r:id'))
        try:
            prs.part.drop_rel(rId)
        except KeyError:
            pass
        sldIdLst.remove(sldId)


def _new_deck_from_template():
    """Create a fresh output deck = title slide only (drop purpose/sample/blank)."""
    prs = Presentation(str(TEMPLATE))
    _delete_slides(prs, [1, 2, 3])   # keep index 0 (title)
    return prs


# ==================================================== helpers: RAG ovals ========
def _place_rag_ovals(slide, status_by_row):
    """Delete status-column ovals and drop clean ones at each status row's centre."""
    status_tbl = _find_table(slide, "Status")
    if status_tbl is None:
        return
    t_left, t_top = status_tbl.left, status_tbl.top
    tbl = status_tbl.table
    col_w = [tbl.columns[i].width for i in range(len(tbl.columns))]
    row_h = [tbl.rows[i].height for i in range(len(tbl.rows))]
    # centre x of the middle (Status) column
    cx = t_left + col_w[0] + col_w[1] / 2

    # remove existing ovals sitting in the status column (left of ~6in, within table height)
    tbl_bottom = t_top + sum(row_h)
    for shp in list(slide.shapes):
        if shp.shape_type == MSO_SHAPE.OVAL or (shp.name or "").startswith("Oval"):
            if shp.left is not None and shp.left < Emu(int(6 * 914400)) and shp.top is not None \
               and t_top <= shp.top <= tbl_bottom:
                shp._element.getparent().remove(shp._element)

    D = Emu(int(0.25 * 914400))
    # rows 1..6 map to STATUS_ROWS in order
    for ridx, label in enumerate(STATUS_ROWS, start=1):
        if ridx >= len(row_h):
            break
        cy = t_top + sum(row_h[:ridx]) + row_h[ridx] / 2
        rag = status_by_row.get(label, {}).get("rag", "grey")
        rag = RAG_ALIASES.get(str(rag).strip().lower(), "grey")
        ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(cx - D / 2), int(cy - D / 2), int(D), int(D))
        ov.fill.solid(); ov.fill.fore_color.rgb = RAG[rag]
        ov.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); ov.line.width = Pt(0.75)
        ov.shadow.inherit = False


# ==================================================== populate one slide ========
def _ensure_rows(table_shape, needed_data_rows):
    """Ensure the table has at least (1 header + needed) rows by cloning the last row."""
    tbl = table_shape.table
    while len(tbl.rows) - 1 < needed_data_rows:
        last_tr = tbl._tbl.tr_lst[-1]
        new_tr = copy.deepcopy(last_tr)
        for tc in new_tr.findall(qn('a:tc')):
            for t in tc.findall('.//' + qn('a:t')):
                t.text = ""
        tbl._tbl.append(new_tr)


def populate_slide(slide, data):
    # banner / project name (rectangle auto-shape containing the placeholder text)
    product = data.get("product_name") or data.get("project_id", "")
    for sh in slide.shapes:
        if sh.has_text_frame:
            txt = sh.text_frame.text.strip()
            if txt in ("Project Name", "Product Name"):
                sh.text_frame.text = product
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(18); r.font.bold = True
            elif txt == "Date":
                sh.text_frame.text = data.get("date", "")

    # status table notes
    st = _find_table(slide, "Status")
    if st:
        tbl = st.table
        row_label = {}
        for r in range(1, len(tbl.rows)):
            row_label[tbl.cell(r, 0).text.strip()] = r
        for label, vals in (data.get("status") or {}).items():
            r = row_label.get(label)
            if r is not None:
                _set_cell(tbl.cell(r, 2), vals.get("notes", ""), size=9)
    _place_rag_ovals(slide, data.get("status") or {})

    # milestones
    ms = _find_table(slide, "Milestone")
    milestones = data.get("milestones") or []
    if ms and milestones:
        _ensure_rows(ms, len(milestones))
        tbl = ms.table
        for i, m in enumerate(milestones, start=1):
            _set_cell(tbl.cell(i, 0), m.get("milestone", ""), size=9)
            _set_cell(tbl.cell(i, 1), m.get("date", ""), size=9)
            _set_cell(tbl.cell(i, 2), m.get("notes", ""), size=9)

    # PDP checklist
    ck = _find_table(slide, "Agreed-upon Objects")
    checklist = data.get("pdp_checklist") or []
    if ck and checklist:
        tbl = ck.table
        existing = {tbl.cell(r, 0).text.strip().lower(): r for r in range(1, len(tbl.rows))}
        for entry in checklist:
            item = entry.get("item", "")
            r = existing.get(item.strip().lower())
            if r is None:
                _ensure_rows(ck, len(tbl.rows))  # add one row
                r = len(tbl.rows) - 1
                _set_cell(tbl.cell(r, 0), item, size=9)
            _set_cell(tbl.cell(r, 1), entry.get("last_reviewed", ""), size=9)
            _set_cell(tbl.cell(r, 2), entry.get("notes", ""), size=9)


# ==================================================== build (create/append) ====
def _next_versioned(path: Path) -> Path:
    """G1: never overwrite. If the deck path is taken by a *different* deck we are not
    appending to, bump v2/v3. (Normal flow appends into the same file, so this only fires
    when the caller explicitly wants a fresh file.)"""
    if not path.exists():
        return path
    stem, ext = path.stem, path.suffix
    n = 2
    while True:
        cand = path.with_name(f"{stem} v{n}{ext}")
        if not cand.exists():
            return cand
        n += 1


def build(data_path: Path, root: Path, append: bool):
    data = json.loads(data_path.read_text(encoding="utf-8"))
    project_id = data["project_id"]
    gov = root / GOV_FOLDER
    gov.mkdir(parents=True, exist_ok=True)
    out_path = gov / f"{project_id} - Project Governance Scorecard.pptx"

    appended = False
    if out_path.exists() and append:
        prs = Presentation(str(out_path))
        appended = True
    else:
        if out_path.exists() and not append:
            out_path = _next_versioned(out_path)
        prs = _new_deck_from_template()
        # set the title-slide product/date if placeholders exist
        for sh in prs.slides[0].shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() in ("Client Project",):
                sh.text_frame.text = data.get("product_name") or project_id

    slide = _clone_blank_slide(prs)
    populate_slide(slide, data)
    n_slides = len(list(prs.slides))
    prs.save(str(out_path))

    return {
        "output": str(out_path),
        "action": "appended new slide" if appended else "created new deck",
        "slides": n_slides,
        "date": data.get("date", ""),
    }


# ============================================================ cli ===============
def main():
    ap = argparse.ArgumentParser(description="Kaizen Project Governance Scorecard generator")
    sub = ap.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("scan", help="scan PDP folders and emit pre-fill JSON")
    s.add_argument("--root", required=True)
    s.add_argument("--project-id", required=True)
    s.add_argument("--out", required=True, help="where to write scan.json")

    b = sub.add_parser("build", help="populate & save/append the scorecard deck")
    b.add_argument("--data", required=True, help="data JSON path")
    b.add_argument("--root", required=True)
    b.add_argument("--no-append", action="store_true", help="force a new deck instead of appending")

    a = ap.parse_args()
    if a.cmd == "scan":
        root = Path(a.root)
        out_deck = root / GOV_FOLDER / f"{a.project_id} - Project Governance Scorecard.pptx"
        res = scan(root, a.project_id, out_deck)
        Path(a.out).write_text(json.dumps(res, indent=2), encoding="utf-8")
        print(json.dumps({"gaps": res["gaps"], "deck_exists": res["deck_exists"],
                          "checklist_found": [c["item"] for c in res["checklist"] if c["found"]],
                          "milestone_candidates": len(res["milestone_candidates"]),
                          "scan_json": a.out}, indent=2))
    elif a.cmd == "build":
        res = build(Path(a.data), Path(a.root), append=not a.no_append)
        print(json.dumps(res, indent=2))


if __name__ == "__main__":
    main()
