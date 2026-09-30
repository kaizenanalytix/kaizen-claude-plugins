"""kvprep.summary -- the input summary workbook, generated rather than typed.

The client's team maintains ``Ensure Data Review Input Summary Q226.xlsx``:
46 sheets, one per channel, each one a grid of periods by series with a
three-row provenance header, a footer of quarter totals with quarter-on-quarter
and year-on-year, and -- at the bottom, in a cell with no formatting -- the
sentence that is the actual point of the document::

    Walmart seems to have a rise in activities in this qtr compared to
    previous, also the spend has incresed for Walmart only

Everything in that file is assembled by hand every quarter. This module
rebuilds it from the run.

**The layout is theirs, deliberately.** Their team reads this file, and a
better arrangement that they have to relearn is not better. The anatomy
reproduced here, read off the Q226 workbook:

=========================  ===================================================
Row 1                      what this sheet was updated for, this quarter
Row 2  ``File Name``       the extract each column came from
Row 3  ``Unit``            Shipments / Bottles / Impressions / Spend / ...
Row 4  (optional group)    DMwC / DMwS / PAB -- present only where they use it
Row 5  ``Channel``         the series name
Rows 6+                    the grid: months (HCP) or week-ending dates (media)
Footer                     YoY, QoQ, then one row per quarter total
Below that                 derived rates (CPM), then the notes and questions
=========================  ===================================================

Monthly sheets carry ``key | Year | Month`` down the left; weekly sheets carry
``index | week ending``. Green fill marks what is new or restated in this
quarter's file -- their convention, and the fastest thing in the document to
read.

**The figures are the client's own, untransformed.** The walkthrough was
explicit: *"if we do any processing on top of these numbers for modelling
purpose, we do not show the processed numbers here. We only show the raw
inputs so that client can validate that this is the correct input that they
have shared with us."* So an HCP sheet shows the monthly figures as they
arrived -- before the redemption lag, before the coupons-per-shipment
multiplier, before the 4.33 divisor -- and the modelled series it becomes is
shown separately, on its own sheet, where nobody can mistake one for the other.

**Three kinds of sheet, never mixed up.** ``generated`` came from a registry
channel and is reproducible. ``inferred`` came from the generic reader, which
guessed the layout of a file nothing in the registry describes. ``carried``
was copied from the previous workbook untouched because this run had no raw
file for it. Every sheet says which it is, in row 1, in words.
"""

from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import yaml

from . import eda

WEEK_START = "Week Starting Date"
WEEK_END = "Week Ending Date"

GENERATED = "generated"
INFERRED = "inferred"
CARRIED = "carried"

ORIGIN_NOTE = {
    GENERATED: (
        "Generated from the raw extract listed against each column. Every figure "
        "traces to a cell in a file you sent us."
    ),
    INFERRED: (
        "Read from a file this pipeline has no configuration for. The layout was "
        "inferred, so treat the column grouping as a proposal and tell us where it "
        "is wrong."
    ),
    CARRIED: (
        "Carried over unchanged from the previous input summary -- no raw file for "
        "this channel was supplied with this run, so nothing here has been "
        "re-derived or re-checked."
    ),
}


class SummaryError(RuntimeError):
    pass


# --------------------------------------------------------------------------
# the spec
# --------------------------------------------------------------------------


@dataclass
class SheetSpec:
    """How one sheet of their workbook is laid out and what feeds it."""

    sheet: str
    channels: list[str] = field(default_factory=list)
    grain: str = "monthly"
    group_row: bool = False
    note: str = ""
    rates: list[dict] = field(default_factory=list)
    order: int = 999
    # How finely the sheet breaks the channel down. This is a summary
    # decision, not a load decision: the OLA load carries one row per raw row
    # across six dimensions, and grouping the summary the same way gives a
    # sheet 180 columns wide that nobody can read. Their own sheet shows
    # eight publishers. Left empty, the channel's own dimensions are used.
    by: list[str] = field(default_factory=list)
    max_series: int = 40


@dataclass
class SummarySpec:
    """``registry/clients/<id>/summary.yaml``, parsed and checked."""

    title: str
    sheets: list[SheetSpec]
    cover_sheet: str = "Model & Hypothesis"
    change_log_sheet: str = "Change Log"
    thresholds: dict = field(default_factory=dict)
    breakdowns: dict = field(default_factory=dict)
    considerations: list[str] = field(default_factory=list)

    def for_channel(self, channel: str) -> SheetSpec | None:
        for s in self.sheets:
            if channel in s.channels:
                return s
        return None


_SPEC_KEYS = {
    "schema_version",
    "title",
    "cover_sheet",
    "change_log_sheet",
    "thresholds",
    "sheets",
    "breakdowns",
    "considerations",
}
_SHEET_KEYS = {"sheet", "channels", "grain", "group_row", "note", "rates", "order",
               "by", "max_series"}
_RATE_KEYS = {"name", "numerator", "denominator", "scale", "format"}


def load_spec(path: str | Path) -> SummarySpec:
    """Parse ``summary.yaml``, refusing anything it does not understand.

    An unknown key here is a typo that would otherwise be ignored silently and
    produce a workbook missing a sheet nobody noticed was configured.
    """
    p = Path(path)
    doc = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    unknown = set(doc) - _SPEC_KEYS
    if unknown:
        raise SummaryError(f"{p}: unknown key(s) {sorted(unknown)}; expected {sorted(_SPEC_KEYS)}")
    sheets = []
    seen: set[str] = set()
    for i, entry in enumerate(doc.get("sheets") or []):
        bad = set(entry) - _SHEET_KEYS
        if bad:
            raise SummaryError(f"{p}: sheets[{i}] has unknown key(s) {sorted(bad)}")
        name = entry.get("sheet")
        if not name:
            raise SummaryError(f"{p}: sheets[{i}] has no 'sheet' name")
        if name in seen:
            raise SummaryError(
                f"{p}: two entries both write the sheet {name!r}. One would silently "
                "overwrite the other."
            )
        seen.add(name)
        grain = entry.get("grain", "monthly")
        if grain not in {"monthly", "weekly"}:
            raise SummaryError(
                f"{p}: sheets[{i}] grain is {grain!r}; expected 'monthly' (the client "
                "sends monthly figures) or 'weekly'."
            )
        for j, r in enumerate(entry.get("rates") or []):
            bad = set(r) - _RATE_KEYS
            if bad:
                raise SummaryError(f"{p}: sheets[{i}].rates[{j}] unknown key(s) {sorted(bad)}")
            for req in ("name", "numerator", "denominator"):
                if not r.get(req):
                    raise SummaryError(f"{p}: sheets[{i}].rates[{j}] has no {req!r}")
        sheets.append(
            SheetSpec(
                sheet=name,
                channels=list(entry.get("channels") or []),
                grain=grain,
                group_row=bool(entry.get("group_row", False)),
                note=entry.get("note", ""),
                rates=list(entry.get("rates") or []),
                order=int(entry.get("order", i)),
                by=list(entry.get("by") or []),
                max_series=int(entry.get("max_series", 40)),
            )
        )
    sheets.sort(key=lambda s: s.order)
    return SummarySpec(
        title=doc.get("title", "Data Review Input Summary"),
        sheets=sheets,
        cover_sheet=doc.get("cover_sheet", "Model & Hypothesis"),
        change_log_sheet=doc.get("change_log_sheet", "Change Log"),
        thresholds=dict(doc.get("thresholds") or {}),
        breakdowns=dict(doc.get("breakdowns") or {}),
        considerations=list(doc.get("considerations") or []),
    )


def spec_path(registry_root: str | Path, client_id: str) -> Path:
    return Path(registry_root) / "clients" / client_id / "summary.yaml"


# --------------------------------------------------------------------------
# reading the previous workbook
# --------------------------------------------------------------------------


@dataclass
class PriorSheet:
    name: str
    grain: str
    index: pd.DatetimeIndex
    columns: dict  # (unit, label) -> pd.Series
    files: dict  # (unit, label) -> source file string
    notes: list[str]
    header_row: int
    data_row: int


@dataclass
class PriorWorkbook:
    path: Path
    sheets: dict  # name -> PriorSheet
    change_log: list  # [(bpm, text)]
    considerations: list
    breakdowns: dict

    def series(self, sheet: str) -> dict:
        s = self.sheets.get(sheet)
        return s.columns if s else {}


_MONTHS = {
    m: i + 1
    for i, m in enumerate(
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    )
}


def _is_year(v) -> bool:
    """A four-digit year, whether the cell holds a number or the text of one."""
    if v is None or isinstance(v, bool):
        return False
    try:
        n = float(str(v).strip())
    except (TypeError, ValueError):
        return False
    return 1990 < n < 2100 and n == int(n)


def _find_label_row(ws, label: str, limit: int = 10) -> tuple[int, int] | None:
    for r in range(1, min(ws.max_row, limit) + 1):
        for c in range(1, min(ws.max_column, 8) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower() == label.lower():
                return r, c
    return None


def read_prior(path: str | Path, *, max_sheets: int | None = None) -> PriorWorkbook:
    """Read the previous input summary: its series, its notes, its change log.

    Read for three reasons, all of which need the actual numbers rather than
    the file as an opaque blob: to compare this quarter's figures against last
    quarter's and report restatements; to carry forward the notes and answers
    already recorded against a sheet; and to keep the change log accreting
    rather than starting again.
    """
    import openpyxl

    p = Path(path)
    with warnings.catch_warnings():
        # Their workbook carries a handful of cells whose serial value is
        # outside Excel's date range and two extensions openpyxl drops. Neither
        # affects the numbers being read, and the warnings would bury the
        # output that matters.
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(p, data_only=True, read_only=False)

    sheets: dict[str, PriorSheet] = {}
    change_log: list[tuple[str, str]] = []
    considerations: list[str] = []
    breakdowns: dict[str, list[str]] = {}

    for ws in wb.worksheets[: max_sheets or len(wb.worksheets)]:
        title = ws.title
        low = title.lower()
        if "change log" in low:
            for r in range(2, ws.max_row + 1):
                bpm = ws.cell(r, 1).value
                txt = ws.cell(r, 2).value
                if isinstance(txt, str) and txt.strip():
                    change_log.append((str(bpm or "").strip(), txt.strip()))
            continue
        if "hypothesis" in low or "model &" in low:
            for r in range(2, ws.max_row + 1):
                ch = ws.cell(r, 1).value
                if not isinstance(ch, str) or not ch.strip():
                    continue
                rest = [
                    str(ws.cell(r, c).value).strip()
                    for c in range(2, ws.max_column + 1)
                    if ws.cell(r, c).value
                ]
                if rest:
                    breakdowns[ch.strip()] = rest
                elif re.match(r"^\s*\d+\.", ch):
                    considerations.append(ch.strip())
            continue

        got = _find_label_row(ws, "Channel") or _find_label_row(ws, "Unit")
        if not got:
            continue
        hdr_row, hdr_col = got
        unit_pos = _find_label_row(ws, "Unit")
        file_pos = _find_label_row(ws, "File Name")
        data_row = hdr_row + 1
        first_data_col = hdr_col + 1

        # The time axis: either a date column or Year + Month columns.
        index_vals: list[pd.Timestamp] = []
        rows: list[int] = []
        grain = "weekly"
        date_col = None
        # Only the label columns are candidates for the time axis -- never the
        # first data column. HCP Coupons' first data column holds cells Excel
        # typed as times, and probing it finds a "date axis" that is really
        # somebody's coupon counts.
        for c in range(1, first_data_col):
            hits = sum(
                1
                for r in range(data_row, min(ws.max_row, data_row + 60) + 1)
                if hasattr(ws.cell(r, c).value, "year")
            )
            if hits > 20:
                date_col = c
                break
        if date_col is not None:
            for r in range(data_row, ws.max_row + 1):
                v = ws.cell(r, date_col).value
                if v is None or not hasattr(v, "year"):
                    continue
                index_vals.append(pd.Timestamp(v))
                rows.append(r)
            if len(index_vals) > 2:
                deltas = pd.Series(index_vals).diff().dt.days.dropna()
                grain = "weekly" if deltas.median() <= 10 else "monthly"
        else:
            ycol = mcol = None
            for c in range(1, first_data_col + 1):
                vals = [ws.cell(r, c).value for r in range(data_row, min(ws.max_row, data_row + 40))]
                # A year may be stored as text. Their Direct Mail sheet stores
                # '2015' as a string, and a numeric-only test skips the whole
                # sheet -- silently, and the sheet then reads as "not in their
                # workbook", which is a much more confusing thing to be told.
                if sum(1 for v in vals if _is_year(v)) > 10:
                    ycol = c
                if sum(1 for v in vals if isinstance(v, str) and v.strip()[:3] in _MONTHS) > 10:
                    mcol = c
            if ycol is None or mcol is None:
                continue
            grain = "monthly"
            year = None
            for r in range(data_row, ws.max_row + 1):
                y = ws.cell(r, ycol).value
                m = ws.cell(r, mcol).value
                if _is_year(y):
                    year = int(float(str(y).strip()))
                if year is None or not isinstance(m, str) or m.strip()[:3] not in _MONTHS:
                    continue
                index_vals.append(pd.Timestamp(year, _MONTHS[m.strip()[:3]], 1))
                rows.append(r)

        if len(index_vals) < 6:
            continue

        idx = pd.DatetimeIndex(index_vals)
        columns: dict[tuple[str, str], pd.Series] = {}
        files: dict[tuple[str, str], str] = {}
        for c in range(first_data_col, ws.max_column + 1):
            label = ws.cell(hdr_row, c).value
            if label is None or str(label).strip() == "":
                continue
            unit = ws.cell(unit_pos[0], c).value if unit_pos else ""
            src = ws.cell(file_pos[0], c).value if file_pos else ""
            vals = []
            for r in rows:
                v = ws.cell(r, c).value
                vals.append(float(v) if isinstance(v, (int, float)) else np.nan)
            s = pd.Series(vals, index=idx, dtype="float64")
            if s.notna().sum() == 0:
                continue
            key = (str(unit or "").strip(), str(label).strip())
            n = 2
            while key in columns:  # repeated (unit, label) pairs do occur
                key = (key[0], f"{str(label).strip()} [{n}]")
                n += 1
            columns[key] = s
            files[key] = str(src or "").strip()

        notes = []
        last_data_row = rows[-1] if rows else data_row
        for r in range(last_data_row, ws.max_row + 1):
            for c in range(1, min(ws.max_column, 8) + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and len(v.strip()) > 40:
                    notes.append(v.strip())
        sheets[title] = PriorSheet(
            name=title,
            grain=grain,
            index=idx,
            columns=columns,
            files=files,
            notes=notes[:40],
            header_row=hdr_row,
            data_row=data_row,
        )

    return PriorWorkbook(
        path=p,
        sheets=sheets,
        change_log=change_log,
        considerations=considerations,
        breakdowns=breakdowns,
    )


# --------------------------------------------------------------------------
# the model
# --------------------------------------------------------------------------


@dataclass
class SheetColumn:
    label: str
    unit: str
    values: pd.Series
    source_file: str = ""
    source_detail: str = ""
    group: str = ""
    new_from: pd.Timestamp | None = None
    transformed_note: str = ""

    # The chart helpers in review_html take a column with these three, so a
    # sheet column can be drawn by exactly the same code as a review column.
    # Two panel renderers drifting apart is how the same series ends up
    # looking like two different series in two documents.
    @property
    def total(self) -> float:
        return float(pd.to_numeric(self.values, errors="coerce").sum())

    @property
    def active_weeks(self) -> int:
        v = pd.to_numeric(self.values, errors="coerce").fillna(0)
        return int((v != 0).sum())

    @property
    def missing_weeks(self) -> int:
        return int(pd.to_numeric(self.values, errors="coerce").isna().sum())


#: Which catalogue group an EDA observation belongs in. The rule suite and the
#: EDA pass ask different questions about the same data, and a reader should
#: not have to know which engine produced a finding in order to find it. One
#: set of groups, both sources, sorted together.
EDA_GROUP_KEY = {
    "Coverage": "completeness",
    "Consistency": "consistency",
    "Movement": "movement",
    "Outliers": "movement",
    "Seasonality": "seasonality",
    "Reconciliation": "reconciliation",
}


@dataclass
class SummarySheet:
    name: str
    grain: str
    origin: str
    index: pd.DatetimeIndex
    columns: list[SheetColumn]
    channels: list[str] = field(default_factory=list)
    #: Results of the KV-C rule suite for the channels feeding this sheet,
    #: already translated by the check catalogue.
    findings: list = field(default_factory=list)
    gate: str = ""
    updated_for: str = ""
    group_row: bool = False
    rates: list[dict] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    carried_notes: list[str] = field(default_factory=list)
    analysis: "eda.ChannelEDA | None" = None
    periods: pd.Series | None = None
    modelled: "pd.DataFrame | None" = None

    @property
    def period_order(self) -> list:
        if self.periods is None:
            return []
        return list(dict.fromkeys(self.periods.tolist()))

    def period_totals(self) -> pd.DataFrame:
        if self.periods is None or not self.columns:
            return pd.DataFrame()
        data = {}
        for c in self.columns:
            v = pd.to_numeric(c.values, errors="coerce").fillna(0.0)
            data[(c.unit, c.label)] = v.groupby(self.periods.values).sum()
        df = pd.DataFrame(data)
        return df.reindex([p for p in self.period_order if p in df.index])


@dataclass
class InputSummary:
    client_id: str
    client_name: str
    load_label: str
    generated: str
    sheets: list[SummarySheet]
    change_log: list = field(default_factory=list)
    breakdowns: dict = field(default_factory=dict)
    considerations: list = field(default_factory=list)
    prior_path: str = ""
    source_files: list = field(default_factory=list)
    #: The check catalogue's group definitions, so the document can title and
    #: describe a group without knowing what is in it.
    groups: dict = field(default_factory=dict)
    gate: str = "CLEAR"

    def all_findings(self) -> list:
        return [f for sh in self.sheets for f in sh.findings]

    def rule_asks(self) -> list:
        """Rule failures that need something from the client, deduped by ask."""
        out: dict[str, dict] = {}
        for sh in self.sheets:
            for f in sh.findings:
                if not f.needs_attention or f.internal_only or not f.ask:
                    continue
                if f.ask.lower().startswith("none"):
                    continue
                e = out.setdefault(
                    f.ask,
                    {"ask": f.ask, "title": f.title, "sheets": [], "tone": f.tone,
                     "rule_id": f.rule_id},
                )
                if sh.name not in e["sheets"]:
                    e["sheets"].append(sh.name)
        return list(out.values())

    @property
    def generated_sheets(self) -> list[SummarySheet]:
        return [s for s in self.sheets if s.origin == GENERATED]

    def observations(self, kinds: set | None = None) -> list:
        out = []
        for s in self.sheets:
            if s.analysis:
                for o in s.analysis.observations:
                    if kinds is None or o.kind in kinds:
                        out.append((s, o))
        return out

    def asks(self) -> list:
        return [(s, o) for s, o in self.observations() if o.severity == "ASK"]


# --------------------------------------------------------------------------
# assembly
# --------------------------------------------------------------------------


def _fiscal_label(idx: pd.DatetimeIndex, year_start: str | None) -> pd.Series:
    from .review import fiscal_periods

    return fiscal_periods(idx, year_start)


def _raw_columns(res, ch_cfg) -> tuple[list[SheetColumn], pd.DatetimeIndex]:
    """One column per (series, measure) from the client's own monthly figures."""
    raw = getattr(res, "raw_monthly", None)
    if raw is None or raw.empty:
        return [], pd.DatetimeIndex([])
    work = raw.copy()
    work["period"] = pd.to_datetime(work["period"], errors="coerce")
    work = work.dropna(subset=["period"])
    idx = pd.DatetimeIndex(sorted(work["period"].unique()))
    # Keyed by (series, metric), not by series. A series contributes one
    # column per measure and each comes from a different raw column; keying on
    # the series alone gives every measure whichever provenance entry happened
    # to be written last, so the Shipments column would name the Bottles column
    # as its source. A provenance row that points at the wrong column is worse
    # than none, because it is checkable and wrong.
    prov = {
        (p.get("series"), p.get("metric")): p
        for p in (res.manifest.get("series_provenance") or [])
    }
    dim_cols = [
        c
        for c in work.columns
        if c not in {"period", "__series"} and c not in set(ch_cfg.metrics)
    ]
    cols: list[SheetColumn] = []
    for sidx, part in work.groupby("__series", sort=True):
        dims = [str(part[c].iloc[0]) for c in dim_cols if part[c].notna().any()]
        base = " · ".join(d for d in dims if d and d.lower() != "nan")
        # A measure that is a registry constant has no raw column and so no
        # heading of its own. Borrowing a sibling measure's heading keeps the
        # two rows of one series under one name; falling back to the dimension
        # path instead makes them look like two unrelated series on the sheet.
        sibling = next(
            (
                prov[(sidx, m)].get("label")
                for m in ch_cfg.metrics
                if (sidx, m) in prov and prov[(sidx, m)].get("label")
            ),
            None,
        )
        for metric in ch_cfg.metrics:
            if metric not in part.columns:
                continue
            p = prov.get((sidx, metric), {})
            s = part.set_index("period")[metric]
            s = pd.to_numeric(s, errors="coerce").groupby(level=0).sum().reindex(idx)
            if s.notna().sum() == 0:
                continue
            label = p.get("label") or sibling or base or metric
            note = ""
            if p.get("transform"):
                mult = float(p.get("multiplier", 1.0) or 1.0)
                bits = []
                if abs(mult - 1.0) > 1e-9:
                    bits.append(f"x {mult:,.2f} coupons per shipment")
                bits.append(f"{p.get('profile')} redemption lag")
                note = (
                    "Shown as received. For the model this becomes "
                    + ", then ".join(bits)
                    + " -- see the Modelled series sheet."
                )
            cols.append(
                SheetColumn(
                    label=str(label),
                    unit=metric,
                    values=s,
                    source_file=Path(str(p.get("file", ""))).name,
                    # The raw heading is kept here even though the column
                    # label is tidied, so a reader can still find the column
                    # in the file it came from.
                    source_detail=" · ".join(
                        str(x) for x in (p.get("sheet"), p.get("block"), label) if x
                    ),
                    group=base if base and base != str(label) else "",
                    transformed_note=note,
                )
            )
    _tidy_labels(cols)
    return cols, idx


def _tidy_labels(cols: list[SheetColumn]) -> None:
    """Make the raw column names readable without losing what identifies them.

    The extract's own headings are written for the person who built the
    extract: ``ENS DMwS TOTAL - BAR (Bottles)``. Three things are wrong with
    that on a client-facing sheet. ``ENS`` is on every column, so it carries no
    information and costs width. ``(Bottles)`` duplicates the Unit row directly
    above it. And the part that actually identifies the series -- ``BAR`` -- is
    the last thing the eye reaches.

    Only the redundant parts are removed. The raw heading stays in the
    provenance row, so nothing is lost: a reader who wants to find the column
    in the source file still can.
    """
    if not cols:
        return
    # Any measure name on this sheet, not just this column's own: a label
    # borrowed from a sibling measure (because this one is a registry
    # constant with no raw column) arrives carrying the sibling's unit, and
    # "DMwC TOTAL - ONC (Shipments)" sitting under a Bottles header reads as
    # a mistake.
    units = {str(c.unit).strip().lower().rstrip("s") for c in cols if c.unit}
    for c in cols:
        label = str(c.label).strip()
        m = re.match(r"^(.*?)\s*\(([^()]*)\)\s*$", label)
        if m and m.group(2).strip().lower().rstrip("s") in units and m.group(1).strip():
            label = m.group(1).strip()
        c.label = label
    # A token every column shares says nothing about any of them. Only leading
    # shared tokens are stripped -- a shared word in the middle may still be
    # doing work in the phrase.
    token_sets = [str(c.label).split() for c in cols]
    if len(cols) < 2 or not all(token_sets):
        return
    common = set(token_sets[0])
    for t in token_sets[1:]:
        common &= set(t)
    if not common:
        return

    proposed = []
    for c in cols:
        parts = str(c.label).split()
        while len(parts) > 1 and parts[0] in common:
            parts.pop(0)
        # A leading separator left behind by the strip is noise.
        while parts and parts[0] in {"-", "--", "·", "|", ":"}:
            parts.pop(0)
        proposed.append(" ".join(parts))

    # Applied only if every label survives it as something a reader could act
    # on. One column reading "(Cases)" is worse than every column carrying a
    # redundant prefix: the prefix is merely wide, the empty label is wrong.
    def usable(s: str) -> bool:
        core = re.sub(r"\([^()]*\)", "", s).strip(" -·|:")
        return bool(core) and len(core) >= 2

    if all(usable(p) for p in proposed):
        for c, p in zip(cols, proposed):
            c.label = p


def _weekly_columns(
    res, ch_cfg, by: list[str], max_series: int = 40
) -> tuple[list[SheetColumn], pd.DatetimeIndex]:
    """One column per (breakdown, measure) from a weekly load, plus its history."""
    from .review import _attach_provenance, _series_columns, _shorten_labels

    frames = []
    hist = getattr(res, "history", None)
    load = res.numeric
    if hist is not None and not hist.empty and WEEK_START in hist.columns:
        h = hist.copy()
        h[WEEK_START] = pd.to_datetime(h[WEEK_START], errors="coerce")
        load_weeks = set(pd.to_datetime(load[WEEK_START]).dropna().unique())
        h = h[~pd.to_datetime(h[WEEK_START]).isin(load_weeks)]
        shared = [c for c in load.columns if c in h.columns]
        if not h.empty:
            frames.append(h[shared])
    frames.append(load)
    combined = pd.concat(frames, ignore_index=True)
    combined[WEEK_START] = pd.to_datetime(combined[WEEK_START], errors="coerce")
    weeks = pd.DatetimeIndex(sorted(combined[WEEK_START].dropna().unique()))
    by = [c for c in by if c in combined.columns]
    sc = _series_columns(combined, by, list(ch_cfg.metrics), weeks, max_series)
    _attach_provenance(sc, res, by)
    _shorten_labels(sc)
    cols = [
        SheetColumn(
            label=c.label,
            unit=c.unit,
            values=c.values,
            source_file=c.source_file,
            source_detail=c.source_detail,
        )
        for c in sc
    ]
    return cols, weeks


def match_prior(
    cols: list[SheetColumn], prior_sheet: "PriorSheet | None", *, min_periods: int = 12
) -> dict:
    """Pair each generated column with its counterpart in the previous summary.

    Matched on the numbers, not on the name. Their column headings are written
    by hand and ours come from the extract -- ``DMwS - ONC Hand Raisers``
    against ``ENS DMwS HR - ONC`` -- so a name match finds almost nothing, and
    the restatement check then silently never runs. Worse, it would appear to
    pass.

    Matching on the data instead is also the stronger test: two series that
    agree on forty months of history are the same series whatever they are
    called, and a pair that agrees on thirty-nine of forty is exactly the
    restatement we are looking for. Each prior column is used once, best match
    first, so two of our series cannot both claim the same one.
    """
    out: dict[str, pd.Series] = {}
    if prior_sheet is None or not prior_sheet.columns:
        return out
    scored: list[tuple[float, int, str, tuple]] = []
    for i, col in enumerate(cols):
        ours = pd.to_numeric(col.values, errors="coerce")
        for key, theirs in prior_sheet.columns.items():
            t = pd.to_numeric(theirs, errors="coerce")
            shared = ours.index.intersection(t.index)
            if len(shared) < min_periods:
                continue
            a, b = ours.reindex(shared), t.reindex(shared)
            both = a.notna() & b.notna() & ((a != 0) | (b != 0))
            if int(both.sum()) < min_periods:
                continue
            a2, b2 = a[both], b[both]
            close = float(np.isclose(a2, b2, rtol=1e-4, atol=0.5).mean())
            if close >= 0.5:
                scored.append((close, i, col.label, key))
    scored.sort(key=lambda x: -x[0])
    used_ours: set[int] = set()
    used_theirs: set[tuple] = set()
    for close, i, _, key in scored:
        if i in used_ours or key in used_theirs:
            continue
        used_ours.add(i)
        used_theirs.add(key)
        name = f"{cols[i].label} ({cols[i].unit})" if cols[i].unit else cols[i].label
        out[name] = prior_sheet.columns[key]
    return out


def build_summary(
    results: list,
    *,
    client_cfg,
    spec: SummarySpec,
    registry_root: str | Path | None = None,
    prior: PriorWorkbook | None = None,
    load_label: str = "",
    today: date | None = None,
    inferred: list | None = None,
) -> InputSummary:
    """Assemble the whole document model: the grids, the rules and the EDA.

    One model behind one document. The rule suite answers "is this safe to
    load" and the EDA answers "what is going on in this data"; both are about
    the same figures, and splitting them across two files meant a reader of
    either one was missing half the answer.
    """
    from . import review as review_mod

    catalogue = review_mod.load_catalogue(registry_root) if registry_root else {}
    sheets: list[SummarySheet] = []
    by_sheet: dict[str, list] = {}
    for res in results:
        ss = spec.for_channel(res.channel)
        if ss is None:
            raise SummaryError(
                f"channel {res.channel!r} ran but no sheet in summary.yaml claims it. "
                "A generated channel with nowhere to go would silently vanish from "
                "the workbook; add it to a sheet's 'channels' list."
            )
        by_sheet.setdefault(ss.sheet, []).append((ss, res))

    year_start = None
    for res in results:
        year_start = (res.manifest.get("fiscal") or {}).get("year_start") or year_start

    for ss in spec.sheets:
        entries = by_sheet.get(ss.sheet)
        if not entries:
            continue
        cols: list[SheetColumn] = []
        idx = pd.DatetimeIndex([])
        channels = []
        modelled_parts = []
        findings: list = []
        gates: list[str] = []
        for spec_sheet, res in entries:
            ch_cfg = client_cfg.channel(res.channel)
            channels.append(res.channel)
            gates.append(res.gate)
            if catalogue:
                findings += review_mod.findings_for(
                    res, catalogue, ch_cfg.label or res.channel
                )
            if ss.grain == "monthly":
                c, i = _raw_columns(res, ch_cfg)
                if not c:
                    # A monthly sheet whose channel produced no raw monthly
                    # frame means the pipeline for it is not cross-tab based.
                    # Falling back to the weekly load would put processed
                    # numbers on a sheet that promises raw ones.
                    raise SummaryError(
                        f"{ss.sheet}: channel {res.channel} has no raw monthly figures, "
                        "so a monthly sheet cannot be built from it without showing "
                        "processed numbers instead. Set grain: weekly for this sheet, "
                        "or check that the channel reads a cross-tab."
                    )
            else:
                by = list(ss.by) or list(
                    getattr(ch_cfg, "summary_by", None) or ch_cfg.group_cols or ch_cfg.dim_cols or []
                )
                c, i = _weekly_columns(res, ch_cfg, by, ss.max_series)
            cols += c
            idx = idx.union(i)
            m = res.numeric.copy()
            m["__channel"] = res.channel
            modelled_parts.append(m)

        for c in cols:
            c.values = c.values.reindex(idx)

        periods = _fiscal_label(idx, year_start)
        series = [
            eda.Series(
                channel=ss.sheet,
                label=c.label,
                unit=c.unit,
                values=c.values,
                grain=ss.grain,
                source_file=c.source_file,
                source_detail=c.source_detail,
                transformed=bool(c.transformed_note),
            )
            for c in cols
        ]
        prior_series = match_prior(cols, prior.sheets.get(ss.sheet)) if prior else {}
        analysis = eda.run_channel(
            series,
            periods,
            channel=ss.sheet,
            label=ss.sheet,
            thresholds=spec.thresholds,
            rates=ss.rates,
            prior=prior_series,
        )
        sheets.append(
            SummarySheet(
                name=ss.sheet,
                grain=ss.grain,
                origin=GENERATED,
                index=idx,
                columns=cols,
                channels=channels,
                updated_for=load_label,
                group_row=ss.group_row,
                rates=ss.rates,
                notes=[o.note for o in analysis.observations if o.severity == "ASK"],
                carried_notes=(prior.sheets[ss.sheet].notes if prior and ss.sheet in prior.sheets else []),
                analysis=analysis,
                periods=periods,
                modelled=pd.concat(modelled_parts, ignore_index=True) if modelled_parts else None,
                findings=findings,
                gate=(
                    "BLOCKED" if any(g.startswith("BLOCKED") for g in gates)
                    else "PROCEED WITH WARNINGS" if any(g != "CLEAR" for g in gates)
                    else "CLEAR"
                ),
            )
        )

    for sh in inferred or []:
        sheets.append(sh)

    change_log = list(prior.change_log) if prior else []
    return InputSummary(
        client_id=client_cfg.id,
        client_name=client_cfg.name,
        load_label=load_label,
        generated=(today or date.today()).strftime("%d %B %Y"),
        sheets=sheets,
        change_log=change_log,
        breakdowns=(prior.breakdowns if prior else {}) or spec.breakdowns,
        considerations=(prior.considerations if prior else []) or spec.considerations,
        prior_path=str(prior.path) if prior else "",
        groups=(catalogue or {}).get("groups", {}),
        gate=(
            "BLOCKED" if any(s.gate == "BLOCKED" for s in sheets)
            else "PROCEED WITH WARNINGS"
            if any(s.gate and s.gate != "CLEAR" for s in sheets)
            else "CLEAR"
        ),
        source_files=sorted(
            {
                Path(str(p)).name
                for res in results
                for p in (res.manifest.get("raw_files") or {}).values()
            }
        ),
    )


def proposed_change_log(summary: InputSummary) -> list[tuple[str, str]]:
    """Change-log entries this run would add, for a person to accept or drop.

    Their change log is the memory of the engagement -- why a number moved two
    years ago, which vendor changed, when a divisor changed. It is maintained
    by hand and it is the most valuable text in the workbook.

    Restatements and stops are detected here, so this proposes the entries
    rather than leaving them to be remembered. It does not write them: a change
    log entry is a claim about intent, and only a person knows the intent.
    """
    out: list[tuple[str, str]] = []
    label = summary.load_label or "this load"
    for sheet, o in summary.observations({"restated"}):
        ev = o.evidence
        out.append(
            (
                label,
                f"{sheet.name}: {o.series} restated -- {ev.get('periods_changed')} of "
                f"{ev.get('periods_compared')} shared periods differ from the previous "
                f"input summary. Reason to be confirmed.",
            )
        )
    for sheet, o in summary.observations({"stopped"}):
        out.append((label, f"{sheet.name}: {o.series} reports no activity in {o.period}."))
    for sheet, o in summary.observations({"started"}):
        out.append((label, f"{sheet.name}: {o.series} appears for the first time in {o.period}."))
    return out
