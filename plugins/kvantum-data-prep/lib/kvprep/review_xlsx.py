"""Render a :class:`kvprep.review.Review` as a workbook the client can work in.

The HTML is for reading; this is for the half of the review that happens in
Excel anyway -- filtering a flag list, pasting a column into an email, adding a
note beside a week. It is laid out the way the client's own input summary is
laid out, one sheet per channel, so their team does not have to learn a second
shape.

Period totals, quarter-on-quarter and year-on-year are written as **formulas
over the week rows above them**, not as numbers computed here. Someone will
edit a cell in this file -- correct a week, blank a row they know is wrong --
and the totals have to follow.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .review import Review, _fmt_date, slug

FONT = "Arial"
_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_LBL = Font(name=FONT, size=9, color="52514E")
_BODY = Font(name=FONT, size=10)
_BOLD = Font(name=FONT, size=10, bold=True)
_TITLE = Font(name=FONT, size=14, bold=True)
_FILL_HEAD = PatternFill("solid", fgColor="1C5CAB")
_FILL_SUB = PatternFill("solid", fgColor="E8EFF9")
_FILL_TOTAL = PatternFill("solid", fgColor="F0EFEC")
_TONE_FILL = {
    "critical": PatternFill("solid", fgColor="F7D7D7"),
    "warning": PatternFill("solid", fgColor="FDEFCC"),
    "serious": PatternFill("solid", fgColor="FBE2D6"),
    "good": PatternFill("solid", fgColor="DCF2DC"),
    "muted": PatternFill("solid", fgColor="F0EFEC"),
}
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(bottom=_THIN)
_NUM = "#,##0.00;(#,##0.00);-"
_PCT = "+0.0%;-0.0%;-"


def _sheet(wb: Workbook, title: str):
    name = title[:31]
    n, i = name, 2
    while n in wb.sheetnames:
        n = f"{name[:28]}~{i}"
        i += 1
    return wb.create_sheet(n)


def _write_table(ws, df: pd.DataFrame, row: int, *, widths=None, tone_col=None) -> int:
    """Write a dataframe as a banded table starting at ``row``; return the next row."""
    if df is None or df.empty:
        ws.cell(row=row, column=1, value="(nothing to report)").font = _LBL
        return row + 2
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=row, column=j, value=str(col))
        c.font, c.fill = _HEAD, _FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(df.columns, start=1):
            v = r[col]
            c = ws.cell(row=row + i, column=j, value=(None if pd.isna(v) else v))
            c.font = _BODY
            c.border = _BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, (int, float)):
                c.number_format = "#,##0"
        if tone_col and tone_col in df.columns:
            fill = _TONE_FILL.get(str(r[tone_col]), None)
            if fill:
                ws.cell(row=row + i, column=1).fill = fill
    for j, col in enumerate(df.columns, start=1):
        w = (widths or {}).get(col)
        if w is None:
            w = min(max(len(str(col)) + 2, int(df[col].astype(str).str.len().max() or 8) + 2), 55)
        ws.column_dimensions[get_column_letter(j)].width = w
    return row + len(df) + 2


def _title(ws, text: str, sub: str = "") -> int:
    ws["A1"] = text
    ws["A1"].font = _TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = _LBL
    ws.freeze_panes = "A4"
    return 4


def _overview(wb: Workbook, rv: Review, audience: str) -> None:
    ws = _sheet(wb, "Overview")
    row = _title(
        ws,
        f"{rv.client_name} — data review{(' · ' + rv.load_label) if rv.load_label else ''}",
        f"Prepared {rv.generated} · overall status: {rv.gate}"
        + (" · internal copy" if audience == "internal" else ""),
    )
    ws.cell(row=row, column=1, value="Files received").font = _BOLD
    row += 1
    files = pd.DataFrame(
        [
            {
                "File": f["name"],
                "Size": f["size"],
                "Modified": f["modified"],
                "Used for": ", ".join(f["channels"]),
            }
            for f in rv.files
        ]
    )
    row = _write_table(ws, files, row, widths={"File": 52, "Used for": 30})

    ws.cell(row=row, column=1, value="Channels").font = _BOLD
    row += 1
    chan = pd.DataFrame(
        [
            {
                "Channel": c.label,
                "Rows prepared": c.rows,
                "Period loaded": c.load_span,
                "Shown in summary": c.history_span,
                "Series": len(c.columns),
                "Items to confirm": sum(
                    1
                    for f in c.findings
                    if f.needs_attention and not (audience == "client" and f.internal_only)
                ),
                "Status": c.gate,
            }
            for c in rv.channels
        ]
    )
    row = _write_table(ws, chan, row, widths={"Channel": 26, "Status": 28})
    ws.cell(row=row, column=1, value="How the calendar works").font = _BOLD
    ws.cell(row=row + 1, column=1, value=rv.calendar_note).font = _LBL
    ws.cell(row=row + 1, column=1).alignment = Alignment(wrap_text=True)


def _asks(wb: Workbook, rv: Review, audience: str) -> None:
    ws = _sheet(wb, "What we need from you")
    row = _title(
        ws,
        "What we need from you",
        "One row per open item. Put your answer in the 'Your response' column.",
    )
    rows = rv.asks(audience)
    df = pd.DataFrame(
        [
            {
                "#": i,
                "Priority": {"critical": "Action needed", "warning": "Confirm",
                             "serious": "Accepted"}.get(r["tone"], "Confirm"),
                "Topic": r["title"],
                "Channels": ", ".join(r["channels"]),
                "What we need": r["ask"],
                "Your response": "",
                "tone": r["tone"],
            }
            for i, r in enumerate(rows, start=1)
        ]
    )
    if df.empty:
        df = pd.DataFrame([{"#": "", "Priority": "", "Topic": "Nothing outstanding",
                            "Channels": "", "What we need": "", "Your response": "",
                            "tone": "good"}])
    show = df.drop(columns=["tone"])
    end = _write_table(
        ws, show, row,
        widths={"#": 5, "Priority": 15, "Topic": 38, "Channels": 26,
                "What we need": 70, "Your response": 40},
        tone_col=None,
    )
    for i, tone in enumerate(df["tone"], start=1):
        fill = _TONE_FILL.get(tone)
        if fill:
            ws.cell(row=row + i, column=2).fill = fill
    ws.cell(row=end, column=1, value="The 'Your response' column is yours to fill in.").font = _LBL


def _quality(wb: Workbook, rv: Review, audience: str) -> None:
    ws = _sheet(wb, "Data quality")
    row = _title(ws, "Data quality", "Every check, grouped by what it looks for.")
    recs = []
    for f in rv.findings(audience):
        recs.append(
            {
                "Group": rv.groups.get(f.group, {}).get("title", f.group),
                "Channel": f.channel,
                "Check": f.title,
                "Result": {"PASS": "Passed", "FAIL": "Needs attention",
                           "WAIVED": "Accepted", "SKIP": "Not tested"}.get(f.status, f.status),
                "Detail": f.headline or f.raw_message,
                "What we need from you": (
                    "" if f.status == "PASS" or f.ask.lower().startswith("none") else f.ask
                ),
                **({"Rule": f.rule_id} if audience == "internal" else {}),
                "tone": f.tone,
            }
        )
    df = pd.DataFrame(recs)
    order = {"critical": 0, "warning": 1, "serious": 2, "muted": 3, "good": 4}
    if not df.empty:
        df = df.assign(_o=df["tone"].map(order)).sort_values(
            ["_o", "Group", "Channel"]).drop(columns="_o")
    show = df.drop(columns=["tone"]) if not df.empty else df
    _write_table(
        ws, show, row,
        widths={"Group": 20, "Channel": 18, "Check": 42, "Result": 18,
                "Detail": 60, "What we need from you": 62},
    )
    for i, tone in enumerate(df["tone"] if not df.empty else [], start=1):
        fill = _TONE_FILL.get(tone)
        if fill:
            ws.cell(row=row + i, column=4).fill = fill


def _evidence(wb: Workbook, rv: Review, audience: str) -> None:
    ws = _sheet(wb, "Flagged rows")
    row = _title(
        ws, "Flagged rows",
        "The actual rows behind every flag, so nothing has to be taken on trust.",
    )
    any_written = False
    for f in rv.attention(audience):
        for t in f.tables:
            frame = t["frame"]
            if frame is None or frame.empty:
                continue
            ws.cell(row=row, column=1,
                    value=f"{f.channel} — {f.title}"
                          + (f" · {t['title']}" if t.get("title") else "")).font = _BOLD
            ws.cell(row=row, column=1).fill = _FILL_SUB
            row += 1
            row = _write_table(ws, frame.head(500), row)
            any_written = True
    if not any_written:
        ws.cell(row=row, column=1, value="No flagged rows — every check passed.").font = _LBL


def _channel(wb: Workbook, ch, rv: Review) -> None:
    """One channel sheet, laid out like the client's own input summary."""
    ws = _sheet(wb, ch.label)
    ws["A1"] = f"{ch.label} — input summary"
    ws["A1"].font = _TITLE
    ws["A2"] = (
        f"{ch.rows:,} rows prepared for {ch.load_span}. "
        + (f"History from {ch.history_source}. " if ch.history_source else "")
        + "Blank cell = no row for that week; 0 = reported as zero."
    )
    ws["A2"].font = _LBL

    hdr = 4  # File name / Source column / Unit / Series live on rows 4..7
    labels = ["Source file", "Source column", "Unit", "Series"]
    for i, lab in enumerate(labels):
        c = ws.cell(row=hdr + i, column=1, value=lab)
        c.font = _LBL if i < 3 else _BOLD
    ws.cell(row=hdr + 3, column=2, value="Period").font = _BOLD

    first_col = 3
    for j, col in enumerate(ch.columns):
        cc = first_col + j
        for i, val in enumerate([col.source_file, col.source_detail, col.unit, col.label]):
            c = ws.cell(row=hdr + i, column=cc, value=val)
            c.font = _LBL if i < 3 else _HEAD
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
            if i == 3:
                c.fill = _FILL_HEAD
        ws.column_dimensions[get_column_letter(cc)].width = 16

    data_start = hdr + 4
    periods = []
    prev, pstart = None, data_start
    for i, wk in enumerate(ch.weeks):
        r = data_start + i
        ws.cell(row=r, column=1, value=_fmt_date(wk)).font = _BODY
        p = ch.periods.get(wk, "")
        ws.cell(row=r, column=2, value=p).font = _LBL
        if p != prev:
            if prev is not None:
                periods.append((prev, pstart, r - 1))
            prev, pstart = p, r
        for j, col in enumerate(ch.columns):
            v = col.values.get(wk)
            c = ws.cell(row=r, column=first_col + j,
                        value=(None if pd.isna(v) else float(v)))
            c.font = _BODY
            c.number_format = _NUM
    if prev is not None:
        periods.append((prev, pstart, data_start + len(ch.weeks) - 1))

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 11
    ws.freeze_panes = ws.cell(row=data_start, column=first_col)

    # -- period totals, as formulas over the rows above --------------------
    row = data_start + len(ch.weeks) + 1
    ws.cell(row=row, column=1, value="Period totals").font = _BOLD
    row += 1
    total_rows = {}
    for label, a, b in periods:
        ws.cell(row=row, column=1, value=f"Total {label}").font = _BOLD
        ws.cell(row=row, column=1).fill = _FILL_TOTAL
        for j in range(len(ch.columns)):
            L = get_column_letter(first_col + j)
            c = ws.cell(row=row, column=first_col + j, value=f"=SUM({L}{a}:{L}{b})")
            c.font, c.number_format, c.fill = _BOLD, _NUM, _FILL_TOTAL
        total_rows[label] = row
        row += 1

    labels_in_order = [p[0] for p in periods]
    if len(labels_in_order) >= 2:
        cur, prv = labels_in_order[-1], labels_in_order[-2]
        ws.cell(row=row, column=1, value=f"QoQ ({cur} vs {prv})").font = _BOLD
        for j in range(len(ch.columns)):
            L = get_column_letter(first_col + j)
            c = ws.cell(
                row=row, column=first_col + j,
                value=f"=IFERROR(({L}{total_rows[cur]}-{L}{total_rows[prv]})"
                      f"/{L}{total_rows[prv]},\"\")",
            )
            c.font, c.number_format = _BODY, _PCT
        row += 1
    if len(labels_in_order) >= 5:
        cur, yr = labels_in_order[-1], labels_in_order[-5]
        ws.cell(row=row, column=1, value=f"YoY ({cur} vs {yr})").font = _BOLD
        for j in range(len(ch.columns)):
            L = get_column_letter(first_col + j)
            c = ws.cell(
                row=row, column=first_col + j,
                value=f"=IFERROR(({L}{total_rows[cur]}-{L}{total_rows[yr]})"
                      f"/{L}{total_rows[yr]},\"\")",
            )
            c.font, c.number_format = _BODY, _PCT
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Coverage").font = _BOLD
    _write_table(ws, ch.gaps(), row + 1, widths={"Breakdown": 30, "Measure": 16})


def render_xlsx(rv: Review, out_path: str | Path, *, audience: str = "client") -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb, rv, audience)
    _asks(wb, rv, audience)
    _quality(wb, rv, audience)
    _evidence(wb, rv, audience)
    for ch in rv.channels:
        _channel(wb, ch, rv)
    p = Path(out_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p
